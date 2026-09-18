import io
import email
from email import policy
from email.parser import BytesParser
from pypdf import PdfReader
from docx import Document

def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    """Extract clean text content from PDF, DOCX, TXT, or EML files."""
    filename_lower = filename.lower()
    
    if filename_lower.endswith(".pdf"):
        return _extract_pdf(file_bytes)
    elif filename_lower.endswith(".docx"):
        return _extract_docx(file_bytes)
    elif filename_lower.endswith(".eml"):
        return _extract_eml(file_bytes)
    else:
        # Default text decode with multi-encoding fallback
        for enc in ["utf-8", "latin-1", "cp1252"]:
            try:
                return file_bytes.decode(enc)
            except UnicodeDecodeError:
                continue
        return file_bytes.decode("utf-8", errors="ignore")

def _extract_pdf(file_bytes: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        pages_text = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                pages_text.append(f"--- Page {i+1} ---\n{text}")
        return "\n\n".join(pages_text) if pages_text else "Empty PDF content."
    except Exception as e:
        return f"PDF Extraction Error: {str(e)}"

def _extract_docx(file_bytes: bytes) -> str:
    try:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(paragraphs)
    except Exception as e:
        return f"DOCX Extraction Error: {str(e)}"

def _extract_eml(file_bytes: bytes) -> str:
    try:
        msg = BytesParser(policy=policy.default).parsebytes(file_bytes)
        sender = msg.get("from", "")
        recipient = msg.get("to", "")
        subject = msg.get("subject", "")
        date = msg.get("date", "")
        
        body = ""
        if msg.is_multipart():
            for part in msg.walk():
                content_type = part.get_content_type()
                if content_type == "text/plain":
                    body += part.get_content()
                elif content_type == "text/html" and not body:
                    body += part.get_content()
        else:
            body = msg.get_content()
            
        return f"Email Metadata:\nFrom: {sender}\nTo: {recipient}\nDate: {date}\nSubject: {subject}\n\nEmail Body:\n{body}"
    except Exception as e:
        return f"EML Extraction Error: {str(e)}"
